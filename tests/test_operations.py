import hashlib
import hmac
import json
import sqlite3
import tempfile
import time
import unittest
import uuid
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
import io
from openpyxl import load_workbook
from urllib.parse import urlencode

import db.connection as db_connection
import db.schema as schema
import server
from db.backups import backup_database
from db.inventory import InventoryUnavailableError, adjust_stock_sync
from db.orders import transition_order_status_sync


TEST_TOKEN = "123456:operations-test-token"


def signed_headers(user_id=101, name="Покупатель"):
    pairs = [
        ("auth_date", str(int(time.time()))),
        ("user", json.dumps({"id": user_id, "first_name": name}, ensure_ascii=False, separators=(",", ":"))),
    ]
    data_check_string = "\n".join(f"{key}={value}" for key, value in sorted(pairs))
    secret = hmac.new(b"WebAppData", TEST_TOKEN.encode(), hashlib.sha256).digest()
    signature = hmac.new(secret, data_check_string.encode(), hashlib.sha256).hexdigest()
    return {"X-Telegram-Init-Data": urlencode([*pairs, ("hash", signature)])}


class OperationsTests(unittest.TestCase):
    def setUp(self):
        self._temporary_directory = tempfile.TemporaryDirectory()
        self._backup_temporary_directory = tempfile.TemporaryDirectory()
        self.database_path = Path(self._temporary_directory.name) / "operations.sqlite"
        self.backup_directory = Path(self._backup_temporary_directory.name) / "backups"
        self._patches = ExitStack()
        self._patches.enter_context(patch.object(schema, "DB_PATH", self.database_path))
        self._patches.enter_context(patch.object(db_connection, "DB_PATH", self.database_path))
        self._patches.enter_context(patch.object(server, "BOT_TOKEN", TEST_TOKEN))
        self._patches.enter_context(patch.object(server.settings, "DELIVERY_ENCRYPTION_ACTIVE_KEY_ID", "test"))
        self._patches.enter_context(
            patch.object(
                server.settings,
                "DELIVERY_ENCRYPTION_KEYS_JSON",
                '{"test":"eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHg="}',
            )
        )
        self._patches.enter_context(
            patch("server.send_telegram_with_keyboard", return_value={"ok": True})
        )
        self._patches.enter_context(patch("server.send_telegram_message", return_value={"ok": True}))
        self._patches.enter_context(patch("server.send_stars_invoice", return_value={"ok": True}))
        schema.initialize_database(self.database_path)
        self.client = server.app.test_client()

    def tearDown(self):
        self._patches.close()
        self._temporary_directory.cleanup()
        self._backup_temporary_directory.cleanup()

    def _execute(self, query, params=()):
        connection = schema.connect(self.database_path)
        try:
            connection.execute(query, params)
            connection.commit()
        finally:
            connection.close()

    def _rows(self, query, params=()):
        connection = schema.connect(self.database_path)
        try:
            return connection.execute(query, params).fetchall()
        finally:
            connection.close()

    def _checkout(self, user_id, quantity, *, payment_method="none"):
        self._execute(
            "UPDATE payment_settings SET setting_value = '0' WHERE setting_key = 'payment_enabled'"
        )
        return self.client.post(
            "/order",
            json={
                "cart": [{"id": 1, "quantity": quantity}],
                "cart_revision": 0,
                "checkout_key": str(uuid.uuid4()),
                "payment_method": payment_method,
                "promo_code": "",
                "delivery": {
                    "method": "sdek_pickup",
                    "recipient_name": "Покупатель",
                    "recipient_phone": "+79991234567",
                    "city": "Москва",
                    "pickup_point": "ПВЗ СДЭК 123",
                },
            },
            headers=signed_headers(user_id),
        )

    def test_finite_inventory_reserves_commits_and_restores(self):
        self._execute("UPDATE books SET stock_quantity = 2 WHERE id = 1")
        response = self._checkout(101, 2)
        self.assertEqual(200, response.status_code)
        order_id = response.get_json()["order_id"]
        self.assertEqual([(1, 2, "reserved")], self._rows(
            "SELECT book_id, quantity, state FROM inventory_reservations"
        ))

        blocked = self._checkout(202, 1)
        self.assertEqual(409, blocked.status_code)
        self.assertEqual(1, self._rows("SELECT COUNT(*) FROM orders")[0][0])

        self.assertTrue(transition_order_status_sync(order_id, "confirmed"))
        self.assertEqual(0, self._rows("SELECT stock_quantity FROM books WHERE id = 1")[0][0])
        self.assertEqual("committed", self._rows("SELECT state FROM inventory_reservations")[0][0])

        self.assertTrue(transition_order_status_sync(order_id, "cancelled"))
        self.assertEqual(2, self._rows("SELECT stock_quantity FROM books WHERE id = 1")[0][0])
        self.assertEqual("released", self._rows("SELECT state FROM inventory_reservations")[0][0])

        self.assertTrue(transition_order_status_sync(order_id, "new", expected_statuses=("cancelled",)))
        self.assertEqual("reserved", self._rows("SELECT state FROM inventory_reservations")[0][0])

    def test_my_orders_are_owner_scoped_and_manual_details_stay_private(self):
        self._execute(
            """
            INSERT INTO orders (user_id, user_name, total, status, payment_method, payment_details_json)
            VALUES (101, 'User', 1200, 'awaiting_payment', 'manual', ?)
            """,
            (json.dumps({"card": "2200123412341234", "recipient": "Recipient"}),),
        )
        order_id = self._rows("SELECT id FROM orders")[0][0]
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 1, 'Книга', 1200)",
            (order_id,),
        )
        owner = self.client.get("/api/orders", headers=signed_headers(101))
        self.assertEqual(200, owner.status_code)
        order = owner.get_json()["orders"][0]
        self.assertEqual(order_id, order["id"])
        self.assertNotIn("payment_details_json", order)
        self.assertNotIn("2200123412341234", json.dumps(order))
        self.assertTrue(order["can_contact_support"])
        self.assertTrue(order["can_resend_manual_details"])

        detail = self.client.get(f"/api/orders/{order_id}", headers=signed_headers(101))
        self.assertEqual(200, detail.status_code)
        self.assertEqual("private, no-store", detail.headers["Cache-Control"])
        receipt = detail.get_json()["order"]
        self.assertEqual([{"book_id": 1, "title": "Книга", "price": 1200, "quantity": 1, "line_total": 1200}], receipt["items"])
        self.assertEqual(1200, receipt["items_subtotal"])
        self.assertNotIn("payment_details_json", receipt)
        self.assertEqual(404, self.client.get(f"/api/orders/{order_id}", headers=signed_headers(202)).status_code)

        outsider = self.client.get("/api/orders", headers=signed_headers(202))
        self.assertEqual([], outsider.get_json()["orders"])
        self.assertEqual(404, self.client.post(
            f"/api/orders/{order_id}/manual-details/resend", headers=signed_headers(202)
        ).status_code)

        resent = self.client.post(
            f"/api/orders/{order_id}/manual-details/resend", headers=signed_headers(101)
        )
        self.assertEqual(200, resent.status_code)
        self.assertEqual(429, self.client.post(
            f"/api/orders/{order_id}/manual-details/resend", headers=signed_headers(101)
        ).status_code)

    def test_order_detail_uses_historical_aggregate_and_discount_snapshots(self):
        self._execute(
            """
            INSERT INTO orders (
                user_id, total, status, payment_method, items_subtotal,
                promo_code_snapshot, promo_discount, bonus_discount
            ) VALUES (101, 1450, 'confirmed', 'manual', 2000, 'SAVE25', 400, 150)
            """
        )
        order_id = self._rows("SELECT id FROM orders")[0][0]
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 1, 'Историческое название', 1000)",
            (order_id,),
        )
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 1, 'Историческое название', 1000)",
            (order_id,),
        )
        self._execute("UPDATE books SET title = 'Новое название' WHERE id = 1")

        response = self.client.get(f"/api/orders/{order_id}", headers=signed_headers(101))

        self.assertEqual(200, response.status_code)
        order = response.get_json()["order"]
        self.assertEqual(
            [{"book_id": 1, "title": "Историческое название", "price": 1000, "quantity": 2, "line_total": 2000}],
            order["items"],
        )
        self.assertEqual("SAVE25", order["promo_code_snapshot"])
        self.assertEqual(400, order["promo_discount"])
        self.assertEqual(150, order["bonus_discount"])
        self.assertEqual(550, order["total_discount"])
        self.assertEqual(1450, order["total"])

    def test_order_support_request_is_owner_scoped_and_idempotent(self):
        self._execute(
            """
            INSERT INTO orders (
                user_id, user_name, total, status, payment_method, items_subtotal,
                promo_code_snapshot, promo_discount, bonus_discount, payment_details_json
            ) VALUES (101, 'Owner', 1450, 'confirmed', 'manual', 2000, 'SAVE25', 400, 150, ?)
            """,
            (json.dumps({"card": "secret-card"}),),
        )
        order_id = self._rows("SELECT id FROM orders")[0][0]
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 1, 'Историческая книга', 1000)",
            (order_id,),
        )
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 1, 'Историческая книга', 1000)",
            (order_id,),
        )

        first = self.client.post(f"/api/orders/{order_id}/support", headers=signed_headers(101))
        repeated = self.client.post(f"/api/orders/{order_id}/support", headers=signed_headers(101))

        self.assertEqual(202, first.status_code)
        self.assertEqual("private, no-store", first.headers["Cache-Control"])
        self.assertFalse(first.get_json()["reused"])
        self.assertEqual(202, repeated.status_code)
        self.assertTrue(repeated.get_json()["reused"])
        self.assertEqual(first.get_json()["request_id"], repeated.get_json()["request_id"])
        self.assertEqual(404, self.client.post(f"/api/orders/{order_id}/support", headers=signed_headers(202)).status_code)

        request_row = self._rows("SELECT receipt_json FROM order_support_requests")[0]
        self.assertIn("Историческая книга", request_row[0])
        self.assertNotIn("secret-card", request_row[0])
        history = self._rows("SELECT text FROM support_messages WHERE user_id = 101")[0][0]
        self.assertIn("Обращение по заказу", history)
        self.assertIn("Историческая книга ×2", history)
        self.assertIn("Промокод: SAVE25", history)
        self.assertIn("Общая скидка: −550 ₽", history)
        self.assertEqual(1, self._rows("SELECT is_support_active FROM users WHERE user_id = 101")[0][0])

        self._execute("UPDATE orders SET status = 'completed' WHERE id = ?", (order_id,))
        self.assertFalse(
            self.client.get("/api/orders", headers=signed_headers(101)).get_json()["orders"][0]["can_contact_support"]
        )
        terminal = self.client.post(f"/api/orders/{order_id}/support", headers=signed_headers(101))
        self.assertEqual(409, terminal.status_code)

    def test_sales_export_is_paid_only_and_neutralizes_formula_cells(self):
        self._execute(
            "INSERT INTO orders (user_id, total, status, payment_method, paid_at) VALUES (1, 100, 'paid', 'none', '2026-01-02 00:00:00')"
        )
        paid_order = self._rows("SELECT id FROM orders")[0][0]
        self._execute(
            "INSERT INTO order_items (order_id, book_id, title, price) VALUES (?, 7, '=formula', 100)",
            (paid_order,),
        )
        self._execute("INSERT INTO orders (user_id, total, status) VALUES (2, 100, 'cancelled')")
        response = self.client.get(
            "/api/admin/export/sales?date_from=2026-01-01&date_to=2026-01-02",
            headers=signed_headers(1),
        )
        self.assertEqual(403, response.status_code)
        with patch.object(server, "ADMIN_IDS", [1]):
            response = self.client.get(
                "/api/admin/export/sales?date_from=2026-01-01&date_to=2026-01-02",
                headers=signed_headers(1),
            )
        self.assertEqual(200, response.status_code)
        text = response.get_data(as_text=True)
        self.assertIn("'=formula", text)
        self.assertNotIn("cancelled", text)

    def test_online_backup_is_verified_and_outside_database_directory(self):
        with patch.object(server.settings, "BACKUP_DIR", str(self.backup_directory)):
            target = backup_database(self.database_path, self.backup_directory)
        self.assertTrue(target.exists())
        connection = sqlite3.connect(target)
        try:
            self.assertEqual("ok", connection.execute("PRAGMA quick_check").fetchone()[0])
            self.assertEqual(6, connection.execute("SELECT COUNT(*) FROM books").fetchone()[0])
        finally:
            connection.close()

        with self.assertRaises(ValueError):
            backup_database(self.database_path, self.database_path.parent)

    def test_backup_uses_active_tenant_context_without_legacy_default(self):
        with (
            patch.object(schema, "DB_PATH", None),
            schema.database_context(self.database_path),
            patch.object(server.settings, "BACKUP_DIR", str(self.backup_directory)),
        ):
            target = backup_database(backup_directory=self.backup_directory)
        self.assertTrue(target.is_file())

    def test_inventory_movements_are_lazy_and_cursor_paginated(self):
        self._execute("UPDATE books SET stock_quantity = 0 WHERE id = 1")
        with patch.object(server, "ADMIN_IDS", [101]):
            initial = self.client.get("/api/admin/inventory", headers=signed_headers(101))
            self.assertEqual(200, initial.status_code)
            self.assertNotIn("movements", initial.get_json())
            self.assertEqual(403, self.client.get(
                "/api/admin/inventory/movements", headers=signed_headers(202)
            ).status_code)
            self.assertTrue(adjust_stock_sync(1, 2, 101, "received"))
            self.assertTrue(adjust_stock_sync(1, 1, 101, "return"))
            first = self.client.get(
                "/api/admin/inventory/movements?limit=1", headers=signed_headers(101)
            )
            self.assertEqual(200, first.status_code)
            first_payload = first.get_json()
            self.assertEqual(1, len(first_payload["movements"]))
            second = self.client.get(
                f"/api/admin/inventory/movements?limit=1&before_id={first_payload['next_before_id']}",
                headers=signed_headers(101),
            )
            self.assertEqual(200, second.status_code)
            self.assertLess(
                second.get_json()["movements"][0]["id"], first_payload["movements"][0]["id"]
            )
    def test_action_journal_xlsx_is_role_safe_and_neutralizes_formulas(self):
        self._execute("UPDATE books SET title = '=HYPERLINK(\"https://invalid\")', stock_quantity = 0 WHERE id = 1")
        with patch.object(server, "ADMIN_IDS", [1]):
            self.assertTrue(adjust_stock_sync(1, 1, 1, "received"))
            response = self.client.get(
                "/api/admin/export/action-journal.xlsx", headers=signed_headers(1)
            )
        self.assertEqual(200, response.status_code)
        self.assertEqual("private, no-store", response.headers["Cache-Control"])
        self.assertEqual(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.mimetype,
        )
        workbook = load_workbook(io.BytesIO(response.data), data_only=False)
        try:
            self.assertEqual(["Движения остатков", "Аудит действий"], workbook.sheetnames)
            title = workbook["Движения остатков"]["D2"]
            self.assertEqual("'=HYPERLINK(\"https://invalid\")", title.value)
            self.assertEqual("s", title.data_type)
            self.assertIsInstance(workbook["Движения остатков"]["G2"].value, int)
        finally:
            workbook.close()
        self.assertEqual(
            403,
            self.client.get(
                "/api/admin/export/action-journal.xlsx", headers=signed_headers(202)
            ).status_code,
        )


if __name__ == "__main__":
    unittest.main()
