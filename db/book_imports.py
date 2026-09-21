from __future__ import annotations

import csv
import hashlib
import io
import json
import sqlite3
import uuid
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path

from controlplane.plan_policy import (
    FEATURE_BOOK_IMPORT,
    LIMIT_BOOKS,
    LIMIT_IMPORTS_PER_DAY,
)
from db.audit import append_audit_event
from db.inventory import _record_movement, _reserved_quantity
from db.schema import connect
from runtime.context import maybe_current_tenant_context
from runtime.features import require_feature
from runtime.quota import require_count_quota, reserve_daily_quota


MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1_000
MAX_COLUMNS = 9
MAX_CELL_LENGTH = 10_000
_BATCH_TTL = timedelta(minutes=30)
_HEADERS = (
    "book_id",
    "title",
    "author",
    "description",
    "price_rub",
    "category",
    "stock_mode",
    "stock_quantity",
)


def build_book_import_template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Книги"
    worksheet.append(_HEADERS)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:H1"
    widths = (12, 36, 28, 52, 14, 24, 16, 18)
    hints = (
        "ID существующей книги для обновления; оставьте пустым для новой.",
        "Обязательное название книги.",
        "Автор книги.",
        "Описание книги.",
        "Обязательная цена в рублях, целое число.",
        "Существующая активная категория или пустое значение.",
        "finite для ограниченного остатка или unlimited для неограниченного.",
        "Количество для finite; пусто для unlimited.",
    )
    for index, (header, width, hint) in enumerate(zip(_HEADERS, widths, hints, strict=True), start=1):
        cell = worksheet.cell(1, index, header)
        cell.font = Font(bold=True)
        cell.comment = Comment(hint, "BookApp")
        worksheet.column_dimensions[cell.column_letter].width = width
    for row in range(2, 102):
        worksheet.cell(row, 7).value = None
    stock_mode = DataValidation(type="list", formula1='"finite,unlimited"', allow_blank=False)
    worksheet.add_data_validation(stock_mode)
    stock_mode.add("G2:G101")
    book_id_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="2000000000", allow_blank=True)
    worksheet.add_data_validation(book_id_validation)
    book_id_validation.add("A2:A101")
    price_validation = DataValidation(type="whole", operator="between", formula1="1", formula2="10000000", allow_blank=False)
    worksheet.add_data_validation(price_validation)
    price_validation.add("E2:E101")
    stock_quantity_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="1000000", allow_blank=True)
    worksheet.add_data_validation(stock_quantity_validation)
    stock_quantity_validation.add("H2:H101")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class BookImportError(ValueError):
    pass


def _text(value: object, field: str, *, required: bool = False, maximum: int = 160) -> str:
    if value is None:
        value = ""
    if not isinstance(value, str):
        value = str(value)
    normalized = value.strip()
    if required and not normalized:
        raise BookImportError(f"{field} is required")
    if len(normalized) > maximum:
        raise BookImportError(f"{field} is too long")
    if any(ord(character) < 32 for character in normalized):
        raise BookImportError(f"{field} contains control characters")
    if normalized[:1] in {"=", "+", "@"}:
        raise BookImportError(f"{field} must not start with a spreadsheet formula")
    return normalized


def _integer(value: object, field: str, *, minimum: int, maximum: int) -> int:
    if isinstance(value, bool):
        raise BookImportError(f"{field} must be an integer")
    if isinstance(value, float) and not value.is_integer():
        raise BookImportError(f"{field} must be an integer")
    try:
        normalized = int(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise BookImportError(f"{field} must be an integer") from exc
    if not minimum <= normalized <= maximum:
        raise BookImportError(f"{field} is out of range")
    return normalized


def _read_csv(data: bytes) -> list[tuple[int, dict[str, object]]]:
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BookImportError("CSV must use UTF-8 encoding") from exc
    header_line = content.splitlines()[0] if content.splitlines() else ""
    delimiter = ";" if header_line.count(";") > header_line.count(",") else ","
    reader = csv.DictReader(io.StringIO(content), delimiter=delimiter)
    if not reader.fieldnames:
        raise BookImportError("Import header is required")
    headers = tuple((header or "").strip() for header in reader.fieldnames)
    if headers != _HEADERS or len(set(headers)) != len(headers):
        raise BookImportError("Import columns do not match the required template")
    rows = []
    for line_number, row in enumerate(reader, start=2):
        if line_number > MAX_ROWS + 1:
            raise BookImportError("Import has too many rows")
        if None in row or len(row) > MAX_COLUMNS:
            raise BookImportError(f"Row {line_number} has too many columns")
        rows.append((line_number, row))
    return rows


def _read_xlsx(data: bytes) -> list[tuple[int, dict[str, object]]]:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise BookImportError("XLSX file is invalid")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 1_000 or sum(info.file_size for info in infos) > MAX_UPLOAD_BYTES * 8:
            raise BookImportError("XLSX archive exceeds limits")
        if any(info.compress_size == 0 < info.file_size or info.file_size > info.compress_size * 100 for info in infos):
            raise BookImportError("XLSX archive compression ratio is unsafe")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BookImportError("XLSX import is unavailable") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise BookImportError("XLSX workbook cannot be read") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise BookImportError("XLSX must contain exactly one sheet")
        sheet = workbook.active
        rows = sheet.iter_rows()
        header_cells = next(rows, None)
        if not header_cells:
            raise BookImportError("Import header is required")
        if len(header_cells) != len(_HEADERS):
            raise BookImportError("Import columns do not match the required template")
        headers = tuple(str(cell.value or "").strip() for cell in header_cells)
        if headers != _HEADERS or any(cell.data_type == "f" for cell in header_cells):
            raise BookImportError("Import columns do not match the required template")
        parsed = []
        for line_number, cells in enumerate(rows, start=2):
            if line_number > MAX_ROWS + 1:
                raise BookImportError("Import has too many rows")
            if len(cells) > MAX_COLUMNS:
                raise BookImportError(f"Row {line_number} has too many columns")
            if any(cell.data_type == "f" for cell in cells):
                raise BookImportError(f"Row {line_number} contains a formula")
            values = [cell.value for cell in cells]
            if all(value in (None, "") for value in values):
                continue
            parsed.append((line_number, dict(zip(_HEADERS, values, strict=True))))
        return parsed
    finally:
        workbook.close()


def _parse_rows(filename: str, data: bytes) -> tuple[str, list[tuple[int, dict[str, object]]]]:
    if not filename or len(filename) > 255:
        raise BookImportError("Import filename is invalid")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return "csv", _read_csv(data)
    if suffix == ".xlsx":
        return "xlsx", _read_xlsx(data)
    raise BookImportError("Only CSV and XLSX files are supported")


def _normalize_row(row: dict[str, object]) -> dict:
    errors: list[str] = []
    try:
        raw_book_id = _text(row.get("book_id"), "book_id", maximum=20)
        book_id = _integer(raw_book_id, "book_id", minimum=1, maximum=2_000_000_000) if raw_book_id else None
        title = _text(row.get("title"), "title", required=True)
        author = _text(row.get("author"), "author")
        description = _text(row.get("description"), "description", maximum=MAX_CELL_LENGTH)
        price = _integer(row.get("price_rub"), "price_rub", minimum=1, maximum=10_000_000)
        category = _text(row.get("category"), "category", maximum=120)
        stock_mode = _text(row.get("stock_mode"), "stock_mode", required=True, maximum=16).lower()
        if stock_mode not in {"finite", "unlimited"}:
            raise BookImportError("stock_mode must be finite or unlimited")
        raw_stock = _text(row.get("stock_quantity"), "stock_quantity", maximum=20)
        if stock_mode == "finite":
            stock_quantity = _integer(raw_stock, "stock_quantity", minimum=0, maximum=1_000_000)
        elif raw_stock:
            raise BookImportError("stock_quantity must be empty for unlimited stock")
        else:
            stock_quantity = None
        return {
            "book_id": book_id,
            "title": title,
            "author": author,
            "description": description,
            "price": price,
            "category": category,
            "stock_quantity": stock_quantity,
            "errors": errors,
        }
    except BookImportError as exc:
        return {"errors": [str(exc)]}


def _category_map(database: sqlite3.Connection) -> dict[str, tuple[int, str]]:
    rows = database.execute(
        "SELECT id, name FROM categories WHERE is_active = 1"
    ).fetchall()
    return {str(name).strip().casefold(): (book_id, str(name)) for book_id, name in rows}


def _validate_against_database(
    database: sqlite3.Connection, row: dict, categories: dict[str, tuple[int, str]]
) -> list[str]:
    if row["errors"]:
        return row["errors"]
    errors: list[str] = []
    category = categories.get(row["category"].casefold()) if row["category"] else None
    if row["category"] and not category:
        errors.append("category does not exist or is inactive")
    row["category_id"] = category[0] if category else 0
    row["category_name"] = category[1] if category else ""
    duplicate = database.execute(
        """
        SELECT id FROM books
        WHERE is_active = 1 AND COALESCE(is_archived, 0) = 0
          AND LOWER(TRIM(title)) = LOWER(TRIM(?))
          AND LOWER(TRIM(COALESCE(author, ''))) = LOWER(TRIM(?))
        LIMIT 1
        """,
        (row["title"], row["author"]),
    ).fetchone()
    if duplicate and duplicate[0] != row["book_id"]:
        errors.append("active book with the same title and author already exists")
    if row["book_id"] and not database.execute(
        "SELECT 1 FROM books WHERE id = ?", (row["book_id"],)
    ).fetchone():
        errors.append("book_id does not exist")
    return errors


def preview_book_import_sync(filename: str, data: bytes, actor_user_id: int) -> dict:
    if not isinstance(actor_user_id, int) or actor_user_id <= 0:
        raise BookImportError("Invalid import owner")
    tenant_context = maybe_current_tenant_context()
    if tenant_context is not None:
        require_feature(FEATURE_BOOK_IMPORT, tenant_context)
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise BookImportError("Import file exceeds the 2 MB limit")
    source_format, raw_rows = _parse_rows(filename, data)
    if not raw_rows:
        raise BookImportError("Import contains no rows")
    database = connect()
    try:
        database.execute("BEGIN IMMEDIATE")
        categories = _category_map(database)
        seen: set[tuple[str, str]] = set()
        normalized_rows = []
        for line_number, raw_row in raw_rows:
            row = _normalize_row(raw_row)
            errors = _validate_against_database(database, row, categories) if not row["errors"] else row["errors"]
            if not errors:
                duplicate_key = (row["title"].casefold(), row["author"].casefold())
                if duplicate_key in seen:
                    errors = ["duplicate title and author in import"]
                seen.add(duplicate_key)
            row["errors"] = errors
            normalized_rows.append((line_number, row))
        batch_id = uuid.uuid4().hex
        content_hash = hashlib.sha256(data).hexdigest()
        expires_at = (datetime.now(UTC) + _BATCH_TTL).strftime("%Y-%m-%d %H:%M:%S")
        database.execute(
            """
            INSERT INTO book_import_batches (
                id, actor_user_id, content_hash, source_format, expires_at
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (batch_id, actor_user_id, content_hash, source_format, expires_at),
        )
        database.executemany(
            """
            INSERT INTO book_import_rows (batch_id, line_number, row_json, errors_json)
            VALUES (?, ?, ?, ?)
            """,
            [
                (
                    batch_id,
                    line_number,
                    json.dumps({key: value for key, value in row.items() if key != "errors"}, ensure_ascii=False, separators=(",", ":")),
                    json.dumps(row["errors"], ensure_ascii=False, separators=(",", ":")),
                )
                for line_number, row in normalized_rows
            ],
        )
        report = {
            "rows": len(normalized_rows),
            "valid_rows": sum(not row["errors"] for _, row in normalized_rows),
            "invalid_rows": sum(bool(row["errors"]) for _, row in normalized_rows),
        }
        database.execute(
            "UPDATE book_import_batches SET report_json = ? WHERE id = ?",
            (json.dumps(report, separators=(",", ":")), batch_id),
        )
        database.commit()
        return {
            "batch_id": batch_id,
            "expires_at": expires_at,
            **report,
            "rows_preview": [
                {"line_number": line, "row": row, "errors": row["errors"]}
                for line, row in normalized_rows[:100]
            ],
        }
    except Exception:
        database.rollback()
        raise
    finally:
        database.close()


def commit_book_import_sync(batch_id: str, actor_user_id: int) -> dict:
    if not isinstance(batch_id, str) or not len(batch_id) == 32:
        raise BookImportError("Invalid import batch")
    tenant_context = maybe_current_tenant_context()
    if tenant_context is not None:
        require_feature(FEATURE_BOOK_IMPORT, tenant_context)
    database = connect()
    database.row_factory = sqlite3.Row
    try:
        database.execute("BEGIN IMMEDIATE")
        batch = database.execute(
            "SELECT * FROM book_import_batches WHERE id = ? AND actor_user_id = ?",
            (batch_id, actor_user_id),
        ).fetchone()
        if not batch:
            raise BookImportError("Import batch not found")
        if batch["state"] == "committed":
            database.commit()
            return json.loads(batch["report_json"])
        if batch["state"] != "previewed" or batch["expires_at"] <= datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S"):
            database.execute("UPDATE book_import_batches SET state = 'expired' WHERE id = ?", (batch_id,))
            database.commit()
            raise BookImportError("Import preview has expired")
        staged = database.execute(
            "SELECT line_number, row_json, errors_json FROM book_import_rows WHERE batch_id = ? ORDER BY line_number",
            (batch_id,),
        ).fetchall()
        if any(json.loads(row["errors_json"]) for row in staged):
            raise BookImportError("Import has invalid rows")
        created_rows = sum(
            json.loads(row["row_json"])["book_id"] is None for row in staged
        )
        active_books = database.execute(
            """
            SELECT COUNT(*) FROM books
            WHERE is_active = 1 AND COALESCE(is_archived, 0) = 0
            """
        ).fetchone()[0]
        require_count_quota(
            LIMIT_BOOKS, active_books, created_rows, context=tenant_context
        )
        reserve_daily_quota(
            database,
            LIMIT_IMPORTS_PER_DAY,
            1,
            context=tenant_context,
        )
        categories = _category_map(database)
        created = 0
        updated = 0
        for staged_row in staged:
            row = json.loads(staged_row["row_json"])
            row["errors"] = []
            errors = _validate_against_database(database, row, categories)
            if errors:
                raise BookImportError(f"Row {staged_row['line_number']} changed after preview")
            if row["book_id"] is None:
                cursor = database.execute(
                    """
                    INSERT INTO books (
                        title, price, category, category_id, author, description, stock_quantity
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        row["title"], row["price"], row["category_name"], row["category_id"],
                        row["author"], row["description"], row["stock_quantity"],
                    ),
                )
                book_id = cursor.lastrowid
                if row["stock_quantity"]:
                    _record_movement(
                        database,
                        book_id=book_id,
                        action="opening_balance",
                        stock_delta=row["stock_quantity"],
                        reason="import",
                        source_key=f"import:{batch_id}:{book_id}",
                        was_available=0,
                    )
                created += 1
                continue
            current = database.execute(
                "SELECT stock_quantity FROM books WHERE id = ?", (row["book_id"],)
            ).fetchone()
            if not current:
                raise BookImportError(f"Row {staged_row['line_number']} book is unavailable")
            reserved = _reserved_quantity(database, row["book_id"])
            if row["stock_quantity"] is not None and row["stock_quantity"] < reserved:
                raise BookImportError(f"Row {staged_row['line_number']} stock is below reservations")
            database.execute(
                """
                UPDATE books SET title = ?, price = ?, category = ?, category_id = ?,
                    author = ?, description = ?, stock_quantity = ?
                WHERE id = ?
                """,
                (
                    row["title"], row["price"], row["category_name"], row["category_id"],
                    row["author"], row["description"], row["stock_quantity"], row["book_id"],
                ),
            )
            previous = current["stock_quantity"]
            if previous is not None and row["stock_quantity"] is not None and previous != row["stock_quantity"]:
                _record_movement(
                    database,
                    book_id=row["book_id"],
                    action="manual_adjustment",
                    stock_delta=row["stock_quantity"] - previous,
                    actor_admin_id=actor_user_id,
                    reason="import",
                    source_key=f"import:{batch_id}:{row['book_id']}",
                    was_available=previous - reserved,
                )
            updated += 1
        result = {"batch_id": batch_id, "created": created, "updated": updated}
        database.execute(
            """
            UPDATE book_import_batches
            SET state = 'committed', committed_at = CURRENT_TIMESTAMP, report_json = ?
            WHERE id = ? AND state = 'previewed'
            """,
            (json.dumps(result, separators=(",", ":")), batch_id),
        )
        append_audit_event(
            database,
            actor_user_id=actor_user_id,
            actor_role="owner",
            source="mini_app",
            action="catalog.import.committed",
            entity_type="book_import_batch",
            entity_id=batch_id,
            details={"batch_id": batch_id, "count": created + updated, "import_hash": batch["content_hash"]},
        )
        database.commit()
        return result
    except Exception:
        database.rollback()
        raise
    finally:
        database.close()
